import csv
import io
import json
from abc import ABC, abstractmethod
import openpyxl


class BaseParser(ABC):
    @abstractmethod
    def parse(self, file_bytes: bytes, sheet_name: str | None = None) -> dict:
        """
        Returns standardized dictionary:
        {
           "headers": ["Name", "Email", ...],
           "rows": [{"Name": "John", "Email": "john@example.com"}],
           "sheets": [{"name": "Sheet1", "rows": 100}],
           "selected_sheet": "Sheet1"
        }
        """
        pass


class CSVParser(BaseParser):
    def parse(self, file_bytes: bytes, sheet_name: str | None = None) -> dict:
        text = file_bytes.decode("utf-8-sig", errors="ignore").strip()
        f = io.StringIO(text)
        reader = csv.reader(f)

        try:
            raw_headers = next(reader)
        except StopIteration:
            return {
                "headers": [],
                "rows": [],
                "sheets": [{"name": "default", "rows": 0}],
                "selected_sheet": "default",
            }

        headers = [h.strip() for h in raw_headers if h.strip()]
        rows = []

        for row in reader:
            if not row or not any(c.strip() for c in row):
                continue
            row_dict = {}
            for idx, header in enumerate(headers):
                if idx < len(row):
                    row_dict[header] = row[idx].strip()
                else:
                    row_dict[header] = ""
            rows.append(row_dict)

        return {
            "headers": headers,
            "rows": rows,
            "sheets": [{"name": "default", "rows": len(rows)}],
            "selected_sheet": "default",
        }


class ExcelParser(BaseParser):
    def parse(self, file_bytes: bytes, sheet_name: str | None = None) -> dict:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet_names = wb.sheetnames

        sheets_info = [
            {"name": name, "rows": wb[name].max_row - 1 if wb[name].max_row > 1 else 0}
            for name in sheet_names
        ]

        target_sheet_name = sheet_name if sheet_name in sheet_names else sheet_names[0]
        sheet = wb[target_sheet_name]

        rows_iter = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows_iter)
        except StopIteration:
            return {
                "headers": [],
                "rows": [],
                "sheets": sheets_info,
                "selected_sheet": target_sheet_name,
            }

        headers = [str(h).strip() for h in raw_headers if h is not None and str(h).strip()]
        rows = []

        for row_vals in rows_iter:
            if not row_vals or not any(v is not None and str(v).strip() for v in row_vals):
                continue
            row_dict = {}
            for idx, header in enumerate(headers):
                if idx < len(row_vals):
                    val = row_vals[idx]
                    row_dict[header] = str(val).strip() if val is not None else ""
                else:
                    row_dict[header] = ""
            rows.append(row_dict)

        return {
            "headers": headers,
            "rows": rows,
            "sheets": sheets_info,
            "selected_sheet": target_sheet_name,
        }


class TSVParser(BaseParser):
    def parse(self, file_bytes: bytes, sheet_name: str | None = None) -> dict:
        text = file_bytes.decode("utf-8-sig", errors="ignore").strip()
        f = io.StringIO(text)
        reader = csv.reader(f, delimiter="\t")

        try:
            raw_headers = next(reader)
        except StopIteration:
            return {
                "headers": [],
                "rows": [],
                "sheets": [{"name": "default", "rows": 0}],
                "selected_sheet": "default",
            }

        headers = [h.strip() for h in raw_headers if h.strip()]
        rows = []

        for row in reader:
            if not row or not any(c.strip() for c in row):
                continue
            row_dict = {}
            for idx, header in enumerate(headers):
                if idx < len(row):
                    row_dict[header] = row[idx].strip()
                else:
                    row_dict[header] = ""
            rows.append(row_dict)

        return {
            "headers": headers,
            "rows": rows,
            "sheets": [{"name": "default", "rows": len(rows)}],
            "selected_sheet": "default",
        }


class JSONParser(BaseParser):
    def parse(self, file_bytes: bytes, sheet_name: str | None = None) -> dict:
        text = file_bytes.decode("utf-8-sig", errors="ignore").strip()
        data = json.loads(text)

        if isinstance(data, dict):
            if "leads" in data and isinstance(data["leads"], list):
                data = data["leads"]
            elif "rows" in data and isinstance(data["rows"], list):
                data = data["rows"]
            else:
                data = [data]

        if not isinstance(data, list) or not data:
            return {
                "headers": [],
                "rows": [],
                "sheets": [{"name": "default", "rows": 0}],
                "selected_sheet": "default",
            }

        headers_set = set()
        rows = []
        for item in data:
            if isinstance(item, dict):
                row_dict = {}
                for k, v in item.items():
                    headers_set.add(str(k))
                    row_dict[str(k)] = str(v) if v is not None else ""
                rows.append(row_dict)

        headers = list(headers_set)
        return {
            "headers": headers,
            "rows": rows,
            "sheets": [{"name": "default", "rows": len(rows)}],
            "selected_sheet": "default",
        }


class ParserFactory:
    _parsers = {
        "csv": CSVParser(),
        "xlsx": ExcelParser(),
        "xls": ExcelParser(),
        "tsv": TSVParser(),
        "json": JSONParser(),
    }

    @classmethod
    def get(cls, file_type: str) -> BaseParser:
        clean_type = file_type.lower().strip()
        if clean_type not in cls._parsers:
            raise ValueError(f"Unsupported import format: '{file_type}'")
        return cls._parsers[clean_type]
