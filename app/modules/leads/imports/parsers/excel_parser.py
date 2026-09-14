import io
import openpyxl
from app.modules.leads.imports.parsers.base import BaseParser


class ExcelParser(BaseParser):
    def parse(self, file_bytes: bytes, sheet_name: str | None = None) -> dict:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet_names = wb.sheetnames

        sheets_info = [
            {"name": name, "rows": wb[name].max_row - 1 if wb[name].max_row > 1 else 0}
            for name in sheet_names
        ]

        target_sheet_name = sheet_name if sheet_name in sheet_names else sheet_names[0]
        sheet = wb[target_sheet_name]

        rows_iter = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows_iter)
        except StopIteration:
            return {
                "headers": [],
                "rows": [],
                "sheets": sheets_info,
                "selected_sheet": target_sheet_name,
            }

        headers = [str(h).strip() for h in raw_headers if h is not None and str(h).strip()]
        rows = []

        for row_vals in rows_iter:
            if not row_vals or not any(v is not None and str(v).strip() for v in row_vals):
                continue
            row_dict = {}
            for idx, header in enumerate(headers):
                if idx < len(row_vals):
                    val = row_vals[idx]
                    row_dict[header] = str(val).strip() if val is not None else ""
                else:
                    row_dict[header] = ""
            rows.append(row_dict)

        return {
            "headers": headers,
            "rows": rows,
            "sheets": sheets_info,
            "selected_sheet": target_sheet_name,
        }
